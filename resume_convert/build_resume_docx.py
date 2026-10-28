# -*- coding: utf-8 -*-
"""Rebuild the resume PDF as an editable DOCX, preserving layout and applying requested content edits."""
import re
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

SRC_XML = r"D:\ChatBIAgent-main\resume_convert\source.xml"
OUT = r"D:\ChatBIAgent-main\resume_convert\黄绍杰_简历_agent开发.docx"

NAVY = RGBColor(0x1F, 0x3B, 0x63)
EA_FONT = "Microsoft YaHei"
LATIN_FONT = "Calibri"


def snap_size(raw):
    """Map measured PDF size back to the nearest Word half-point size."""
    if raw < 10:
        return 8.5
    if raw < 12:
        return 11.5
    if raw < 14:
        return 12.0
    return 19.0


def parse_runs(path):
    tree = ET.parse(path)
    root = tree.getroot()
    fonts = {}
    for fs in root.iter("fontspec"):
        fonts[fs.get("id")] = (float(fs.get("size")) / 1.5, fs.get("family"))
    page = root.find("page")
    items = []
    for t in page.findall("text"):
        size, family = fonts[t.get("font")]
        items.append({
            "top": float(t.get("top")) / 1.5,
            "left": float(t.get("left")) / 1.5,
            "width": float(t.get("width")) / 1.5,
            "size": snap_size(size),
            "family": family,
            "bold": t.find("b") is not None,
            "text": "".join(t.itertext()),
        })
    items.sort(key=lambda x: (x["top"], x["left"]))

    lines = []
    for it in items:
        if lines and abs(lines[-1]["top"] - it["top"]) <= 5.5:
            lines[-1]["runs"].append(it)
        else:
            lines.append({"top": it["top"], "runs": [it]})
    paragraphs = []
    for ln in lines:
        runs = sorted(ln["runs"], key=lambda r: r["left"])
        # drop fully empty spacer lines
        if not "".join(r["text"] for r in runs).strip():
            if paragraphs:
                paragraphs[-1]["trailing_space"] = True
            continue
        text = "".join(r["text"] for r in runs)
        first = runs[0]
        start_x = first["left"]
        if text.lstrip().startswith("•"):
            paragraphs.append({"runs": runs, "kind": "bullet"})
        elif (paragraphs and paragraphs[-1]["kind"] == "bullet"
              and abs(start_x - 43.3) < 2.5):
            paragraphs[-1]["runs"].extend(runs)
        elif (paragraphs and abs(start_x - 32.7) < 2.5
              and paragraphs[-1].get("right_edge", 0) > 565
              and paragraphs[-1]["kind"] not in ("bullet",)):
            paragraphs[-1]["runs"].extend(runs)
        else:
            paragraphs.append({"runs": runs, "kind": "body"})
        paragraphs[-1]["right_edge"] = max(r["left"] + r["width"] for r in runs)
        paragraphs[-1]["top"] = ln["top"]
    return paragraphs


def merge_runs(runs):
    """Merge consecutive runs that share identical formatting."""
    out = []
    for r in runs:
        key = (r["bold"], r["size"], r["family"])
        if out and out[-1][0] == key:
            out[-1][1].append(r)
        else:
            out.append([key, [r]])
    merged = []
    for key, group in out:
        merged.append({
            "bold": key[0],
            "size": key[1],
            "family": key[2],
            "text": "".join(g["text"] for g in group),
        })
    return merged


def set_run_font(run, family, size, bold):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = LATIN_FONT if "Calibri" in family else EA_FONT
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), LATIN_FONT if "Calibri" in family else EA_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT if "Calibri" in family else EA_FONT)
    rfonts.set(qn("w:eastAsia"), EA_FONT)
    rfonts.set(qn("w:cs"), LATIN_FONT)


def add_bottom_border(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")  # 1 pt
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3B63")
    pbdr.append(bottom)
    ppr.append(pbdr)


def keep_with_next(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:keepNext")
    ppr.append(el)


def main():
    paragraphs = parse_runs(SRC_XML)
    for p in paragraphs:
        p["runs"] = merge_runs(p["runs"])

    # ------------------------------------------------------------------
    # Apply requested content edits
    # ------------------------------------------------------------------
    def replace_paragraph(pred, new_runs):
        for p in paragraphs:
            if pred(p):
                p["runs"] = new_runs
                return
        raise RuntimeError("target paragraph not found")

    def plain(p):
        return "".join(r["text"] for r in p["runs"])

    # 1. Project intro: single e-commerce dataset (Olist) only
    replace_paragraph(
        lambda p: plain(p).startswith("项目简介：面向业务人员的自然语言查数"),
        [
            {"bold": True, "size": 8.5, "family": "YaHei", "text": "项目简介"},
            {"bold": False, "size": 8.5, "family": "Calibri",
             "text": "：面向业务人员的自然语言查数（ChatBI）系统，覆盖电商（Olist）数据集；"
                     "以“NL → IR → SQL”两阶段编译与 LangGraph 多智能体底座把自然语言转为可信 SQL，"
                     "并提供根因分析、字段级评测与全链路审计。"},
        ],
    )

    # 2. Semantic-layer / SQL Guard bullet: drop the dense parenthetical enumerations
    replace_paragraph(
        lambda p: plain(p).startswith("• 语义层确定性编译与 SQL Guard 三层防护"),
        [
            {"bold": False, "size": 8.5, "family": "Calibri", "text": "• "},
            {"bold": True, "size": 8.5, "family": "YaHei", "text": "语义层确定性编译与 SQL Guard 三层防护"},
            {"bold": False, "size": 8.5, "family": "Calibri",
             "text": "：语义层内置 14 个基础指标、12 个维度与 6 类复合计划，"
                     "显式记录聚合方式、可用维度与 JOIN 关系，"
                     "可信 JSON 注册表整体校验后加载且不允许覆盖既有键，"
                     "已治理指标命中即编译、不经过 LLM 概率生成；"
                     "SQL Guard 用 sqlglot AST 做静态校验 → 执行失败 LLM 自修复 → 结果合理性检查。"},
        ],
    )

    # 3. LLM / Agent skills: add Agent Harness and Skill mechanism
    replace_paragraph(
        lambda p: plain(p).startswith("• 大模型与 Agent"),
        [
            {"bold": False, "size": 8.5, "family": "Calibri", "text": "• "},
            {"bold": True, "size": 8.5, "family": "YaHei", "text": "大模型与 Agent"},
            {"bold": False, "size": 8.5, "family": "Calibri",
             "text": "：Transformer、Prompt 与上下文工程、LangGraph StateGraph 编排、"
                     "多智能体协作模式（Supervisor / ReAct / Coordinator / Reflection）、"
                     "MCP 工具协议、Function Calling、Agent Harness 与 Skill 机制。"},
        ],
    )

    # ------------------------------------------------------------------
    # Build document
    # ------------------------------------------------------------------
    doc = Document()
    section = doc.sections[0]
    section.page_width = Pt(612)
    section.page_height = Pt(792)
    section.left_margin = Pt(32.4)
    section.right_margin = Pt(32.4)
    section.top_margin = Pt(28.35)
    section.bottom_margin = Pt(32.4)
    section.header_distance = Pt(14)
    section.footer_distance = Pt(14)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(8.5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), EA_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(14)

    headings = {"个人背景", "项目经历", "专业技能"}

    for p in paragraphs:
        text = plain(p).strip()
        runs = p["runs"]
        # trim trailing whitespace-only runs
        while runs and not runs[-1]["text"].strip():
            runs.pop()
        if runs:
            runs[-1]["text"] = runs[-1]["text"].rstrip()

        para = doc.add_paragraph()
        pf = para.paragraph_format

        if p is paragraphs[0]:
            # name / title
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(25)
        elif text in headings:
            pf.space_before = Pt(4)
            pf.space_after = Pt(1.5)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(18)
            add_bottom_border(para)
            keep_with_next(para)
        elif re.match(r"^[•]", text):
            pf.left_indent = Pt(10.9)
            pf.first_line_indent = Pt(-10.9)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(14)
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
        elif runs and runs[0]["bold"] and runs[0]["size"] >= 12:
            # project entry title line
            pf.space_before = Pt(8)
            pf.space_after = Pt(2)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(16.5)
            keep_with_next(para)
        elif p is paragraphs[1]:
            pf.space_before = Pt(2.5)
            pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(16)
        elif p is paragraphs[2]:
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(16)
        elif text.startswith(("南开大学",)):
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(14)
            pf.space_before = Pt(0)
            pf.space_after = Pt(1.5)
        else:
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(14)
            pf.space_before = Pt(0)
            pf.space_after = Pt(1.5)

        for r in runs:
            text = r["text"]
            # The source PDF aligns columns with full-width CJK spacing; recreate
            # those gaps with ideographic spaces so the Word layout matches.
            if text and set(text) == {" "} and "YaHei" in r["family"]:
                if r["size"] <= 9:
                    text = "\u3000" * len(text)
                else:
                    text = "\u3000" * max(1, round(len(text) * 7.33 / r["size"]))
            run = para.add_run(text)
            set_run_font(run, r["family"], r["size"], r["bold"])

    # document metadata
    doc.core_properties.title = "黄绍杰 - 软件工程硕士研究生 - 简历"
    doc.core_properties.author = "黄绍杰"

    doc.save(OUT)
    print("saved", OUT)
    print("paragraphs:", len(doc.paragraphs))


if __name__ == "__main__":
    main()
